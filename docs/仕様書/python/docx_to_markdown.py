#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Word文書(.docx)をMarkdownに変換するスクリプト（アウトライン対応版）
"""

import docx
import re
import sys

def convert_docx_to_markdown(docx_file, output_file):
    """Word文書をMarkdownに変換（表と段落を正しい順序で処理）"""

    doc = docx.Document(docx_file)
    output = []

    # 段落と表を混在させて処理するため、document.element.bodyを使用
    for element in doc.element.body:
        # 段落の場合
        if element.tag.endswith('p'):
            # 対応する段落オブジェクトを探す
            for para in doc.paragraphs:
                if para._element == element:
                    text = para.text.strip()
                    style_name = para.style.name

                    if style_name.startswith('Heading'):
                        # 見出しレベルを取得
                        match = re.search(r'Heading (\d+)', style_name)
                        if match:
                            level = int(match.group(1))
                            output.append('#' * level + ' ' + text)
                            output.append('')
                        else:
                            output.append('## ' + text)
                            output.append('')
                    elif style_name == 'Title':
                        output.append('# ' + text)
                        output.append('')
                    elif style_name == 'Subtitle':
                        output.append('*' + text + '*')
                        output.append('')
                    else:
                        # 通常の段落でも、番号付きヘッダーを検出
                        if text:
                            # "1. " "1.1. " "1.1.1. " などの形式を検出
                            heading_match = re.match(r'^(\d+(?:\.\d+)*)\.\s+(.+)$', text)
                            if heading_match:
                                number = heading_match.group(1)
                                title = heading_match.group(2)
                                level = number.count('.') + 2  # "1." -> ##, "1.1." -> ###
                                output.append('#' * level + ' ' + number + '. ' + title)
                                output.append('')
                            else:
                                output.append(text)
                                output.append('')
                        else:
                            output.append('')
                    break

        # 表の場合
        elif element.tag.endswith('tbl'):
            # 対応する表オブジェクトを探す
            for table in doc.tables:
                if table._element == element:
                    # ヘッダー行
                    if table.rows:
                        header_cells = [cell.text.strip() for cell in table.rows[0].cells]
                        output.append('| ' + ' | '.join(header_cells) + ' |')
                        output.append('|' + '|'.join(['---'] * len(header_cells)) + '|')

                        # データ行
                        for row in table.rows[1:]:
                            cells = [cell.text.strip() for cell in row.cells]
                            output.append('| ' + ' | '.join(cells) + ' |')
                        output.append('')
                    break

    # ファイルに書き込み
    with open(output_file, 'w', encoding='utf-8') as f:
        f.write('\n'.join(output))

    print(f'✅ Successfully converted: {output_file}')

if __name__ == '__main__':
    if len(sys.argv) < 3:
        print('Usage: python docx_to_markdown.py <input.docx> <output.md>')
        sys.exit(1)

    input_file = sys.argv[1]
    output_file = sys.argv[2]

    convert_docx_to_markdown(input_file, output_file)
